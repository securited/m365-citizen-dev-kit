# /// script
# requires-python = ">=3.12,<3.13"
# dependencies = [
#   "pandas",
#   "openpyxl",
#   "rich",
# ]
# ///
"""CSV Merge Report — Packaged Python sample (Pattern B — Script).

Merges every CSV file in the input/ folder into one formatted Excel
workbook with a summary sheet, then exits. On first run, when input/ is
empty, it seeds three sample CSVs so the app demonstrates itself.

Demonstrates the Packaged Python pattern:
- PEP 723 inline metadata is the only dependency declaration
- Approved stack only: pandas, openpyxl, rich + stdlib
- Logs under %LOCALAPPDATA%\\csv-merge-report\\ — never system locations
- Single file, no secrets, no network access, runs to completion
"""
from __future__ import annotations

import logging
import os
import re
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from rich.console import Console
from rich.panel import Panel
from rich.table import Table

APP_NAME = "csv-merge-report"
APP_DIR = Path(__file__).resolve().parent
INPUT_DIR = APP_DIR / "input"
OUTPUT_DIR = APP_DIR / "output"

console = Console()


def setup_logging() -> logging.Logger:
    log_dir = Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / APP_NAME
    log_dir.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        filename=log_dir / "run.log",
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
    )
    return logging.getLogger(APP_NAME)


SAMPLE_CSVS = {
    "imaging_orders.csv": (
        "order_id,part_number,description,quantity,unit_price,status\n"
        "SO-310021,PART-77210,X-Ray Tube Housing,1,1825.00,Shipped\n"
        "SO-310044,PART-77342,Collimator Assembly,2,640.50,Delivered\n"
        "SO-310067,PART-77415,Detector Panel,1,3120.00,Open\n"
        "SO-310102,PART-77108,Image Intensifier,1,2275.25,Backordered\n"
        "SO-310145,PART-77342,Collimator Assembly,1,640.50,Delivered\n"
    ),
    "biomed_orders.csv": (
        "order_id,part_number,description,quantity,unit_price,status\n"
        "SO-310203,PART-88110,SpO2 Sensor,6,118.75,Delivered\n"
        "SO-310218,PART-88254,ECG Lead Set,4,86.20,Shipped\n"
        "SO-310240,PART-88307,NIBP Cuff Adult,8,42.10,Delivered\n"
        "SO-310266,PART-88412,Monitor Battery Pack,3,210.00,Open\n"
        "SO-310291,PART-88110,SpO2 Sensor,2,118.75,Open\n"
        "SO-310310,PART-88254,ECG Lead Set,5,86.20,Delivered\n"
    ),
    "lab_orders.csv": (
        "order_id,part_number,description,quantity,unit_price,status\n"
        "SO-310402,PART-99120,Centrifuge Rotor,1,975.00,Shipped\n"
        "SO-310433,PART-99245,Reagent Pump,2,388.40,Delivered\n"
        "SO-310470,PART-99318,Sample Probe,3,154.60,Open\n"
        "SO-310498,PART-99245,Reagent Pump,1,388.40,Backordered\n"
    ),
}


def seed_sample_data(log: logging.Logger) -> None:
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    for name, content in SAMPLE_CSVS.items():
        (INPUT_DIR / name).write_text(content, encoding="utf-8")
    log.info("Seeded %d sample CSV files into %s", len(SAMPLE_CSVS), INPUT_DIR)
    console.print(
        f"[yellow]No CSV files found — created {len(SAMPLE_CSVS)} sample files "
        f"in [bold]{INPUT_DIR}[/bold] so you can see the report format. "
        "Drop your own CSVs there and run again.[/yellow]\n"
    )


def safe_sheet_name(stem: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "_", stem)[:31] or "Sheet"
    candidate = name
    n = 2
    while candidate.lower() in used:
        suffix = f"_{n}"
        candidate = name[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def style_worksheet(ws) -> None:
    bold = Font(bold=True)
    for col in range(1, ws.max_column + 1):
        ws.cell(row=1, column=col).font = bold
        width = len(str(ws.cell(row=1, column=col).value or ""))
        for row in range(2, min(ws.max_row, 200) + 1):
            width = max(width, len(str(ws.cell(row=row, column=col).value or "")))
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> int:
    log = setup_logging()
    log.info("Run started")
    console.print(Panel.fit("[bold]CSV Merge Report[/bold]\nPackaged Python sample — Pattern B (Script)"))

    csv_files = sorted(INPUT_DIR.glob("*.csv")) if INPUT_DIR.exists() else []
    if not csv_files:
        seed_sample_data(log)
        csv_files = sorted(INPUT_DIR.glob("*.csv"))

    loaded: list[tuple[Path, pd.DataFrame]] = []
    for file in csv_files:
        try:
            loaded.append((file, pd.read_csv(file)))
        except Exception as exc:
            log.exception("Failed to read %s", file)
            console.print(f"[red]Skipping {file.name}: {exc}[/red]")

    if not loaded:
        console.print("[red]No readable CSV files — nothing to report.[/red]")
        log.error("Run aborted: no readable input files")
        return 1

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = OUTPUT_DIR / f"merge_report_{stamp}.xlsx"

    summary = pd.DataFrame(
        [{"File": f.name, "Rows": len(df), "Columns": len(df.columns)} for f, df in loaded]
    )

    used_names: set[str] = set()
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name=safe_sheet_name("Summary", used_names), index=False)
        for file, df in loaded:
            df.to_excel(writer, sheet_name=safe_sheet_name(file.stem, used_names), index=False)
        for ws in writer.book.worksheets:
            style_worksheet(ws)

    table = Table(title="Files merged")
    table.add_column("File")
    table.add_column("Rows", justify="right")
    table.add_column("Columns", justify="right")
    for file, df in loaded:
        table.add_row(file.name, str(len(df)), str(len(df.columns)))
    console.print(table)
    console.print(Panel.fit(f"[green]Report written:[/green] [bold]{out_path}[/bold]"))

    log.info("Run finished: %d files merged into %s", len(loaded), out_path)
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:
        logging.getLogger(APP_NAME).exception("Unhandled error")
        console.print(f"[red]Unexpected error: {exc}[/red]")
        sys.exit(1)
